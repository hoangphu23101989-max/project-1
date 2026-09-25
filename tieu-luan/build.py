# -*- coding: utf-8 -*-
"""Dựng tiểu luận Triết học (.docx, .pdf) theo quy định trình bày của môn học.

Quy định áp dụng (theo hướng dẫn của giảng viên phụ trách):
  - Times New Roman; từ "Mục lục" trở đi cỡ 13 (tiêu đề chương cỡ 14).
  - Lề trên 3,5 cm; lề dưới 3,0 cm; lề trái 3,5 cm; lề phải 2,0 cm.
  - Dãn dòng 1,15; số trang ở phía trên, giữa; trang 1 là "Phần mở đầu".

Mục lục được dựng tĩnh qua hai lượt: lượt 1 dựng tài liệu và chuyển sang PDF để
xác định số trang của từng đề mục; lượt 2 điền số trang vào mục lục.

Chạy:  python3 build.py
"""
import os
import re
import subprocess
import sys
import unicodedata

import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import noidung as nd  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
TEN_TEP = "TieuLuan_TrietHoc_NguyenDoanDiemNgoc_25C56077"
FONT = "Times New Roman"
CO_CHU = 13
CO_CHUONG = 14
DAN_DONG = 1.15
BE_RONG_CHU = 21.0 - 3.5 - 2.0  # cm


# ---------------------------------------------------------------- tiện ích XML
def _set_fonts(rpr_parent, size=None, bold=None, italic=None):
    """Đặt phông Times New Roman cho mọi bảng mã (kể cả tiếng Việt) và bỏ phông chủ đề."""
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    return rpr


def style_font(style, size, bold=False, italic=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    _set_fonts(style.element)
    rpr = style.element.get_or_add_rPr()
    color = rpr.find(qn("w:color"))
    if color is not None:
        rpr.remove(color)
    style.font.color.rgb = RGBColor(0, 0, 0)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "vi-VN")


def para_fmt(p, align=None, first=None, left=None, right=None, before=None, after=None,
             spacing=None, keep_next=None, page_break=None):
    f = p.paragraph_format
    if align is not None:
        f.alignment = align
    if first is not None:
        f.first_line_indent = Cm(first)
    if left is not None:
        f.left_indent = Cm(left)
    if right is not None:
        f.right_indent = Cm(right)
    if before is not None:
        f.space_before = Pt(before)
    if after is not None:
        f.space_after = Pt(after)
    if spacing is not None:
        f.line_spacing = spacing
    if keep_next is not None:
        f.keep_with_next = keep_next
    if page_break is not None:
        f.page_break_before = page_break
    return p


TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")

REF_KEYS = [k for k, _ in nd.TAI_LIEU_VIET] + [k for k, _ in nd.TAI_LIEU_ANH]
REF_NUM = {k: i + 1 for i, k in enumerate(REF_KEYS)}
assert len(REF_NUM) == len(REF_KEYS), "Khóa tài liệu tham khảo bị trùng"
CITED = set()
CITE = re.compile(r"\[@([^\]]+)\]")


def cite(text):
    """Thay [@khoa1; @khoa2] bằng số thứ tự [n1], [n2] theo danh mục tài liệu tham khảo."""
    def rep(m):
        nums = []
        for k in m.group(1).split(";"):
            k = k.strip().lstrip("@")
            if k not in REF_NUM:
                raise KeyError("Trích dẫn không có trong danh mục: " + k)
            CITED.add(k)
            nums.append(REF_NUM[k])
        return ", ".join("[%d]" % n for n in sorted(set(nums)))
    return CITE.sub(rep, text)


def add_rich(p, text, size=None, bold=None, italic=None):
    """Thêm văn bản có định dạng nội dòng: **đậm**, *nghiêng*, _{dưới}, ^{trên}."""
    text = cite(text)
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, sub, sup = bold, italic, False, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*") and part.endswith("*"):
            part, i = part[1:-1], True
        elif part.startswith("_{"):
            part, sub = part[2:-1], True
        elif part.startswith("^{"):
            part, sup = part[2:-1], True
        r = p.add_run(part)
        if size is not None:
            r.font.size = Pt(size)
        if b is not None:
            r.font.bold = b
        if i is not None:
            r.font.italic = i
        if sub:
            r.font.subscript = True
        if sup:
            r.font.superscript = True
        _set_fonts(r._r)
    return p


def bottom_rule(p):
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    bdr.append(bottom)
    # pBdr phải đứng trước spacing/ind/jc trong pPr (theo schema OOXML)
    after = [qn("w:" + t) for t in ("shd", "tabs", "spacing", "ind", "jc", "rPr")]
    nxt = next((c for c in ppr if c.tag in after), None)
    if nxt is not None:
        nxt.addprevious(bdr)
    else:
        ppr.append(bdr)


def page_borders(section):
    sect = section._sectPr
    b = OxmlElement("w:pgBorders")
    b.set(qn("w:offsetFrom"), "text")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "thickThinSmallGap")
        e.set(qn("w:sz"), "24")
        e.set(qn("w:space"), "8")
        e.set(qn("w:color"), "auto")
        b.append(e)
    sect.find(qn("w:pgMar")).addnext(b)


def remove_page_borders(section):
    for b in section._sectPr.findall(qn("w:pgBorders")):
        section._sectPr.remove(b)


def page_numbering(section, fmt, start=1):
    sect = section._sectPr
    for old in sect.findall(qn("w:pgNumType")):
        sect.remove(old)
    pn = OxmlElement("w:pgNumType")
    pn.set(qn("w:fmt"), fmt)
    pn.set(qn("w:start"), str(start))
    cols = sect.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pn)
    else:
        sect.append(pn)


def header_page_field(section):
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    para_fmt(hp, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, after=0)
    for kind, val in (("begin", None), ("instr", "PAGE"), ("separate", None),
                      ("text", "1"), ("end", None)):
        r = hp.add_run()
        r.font.size = Pt(CO_CHU)
        _set_fonts(r._r)
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = val
            r._r.append(it)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = val
            r._r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def set_margins(section):
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3.5), Cm(3.0)
    section.left_margin, section.right_margin = Cm(3.5), Cm(2.0)
    section.header_distance, section.footer_distance = Cm(1.8), Cm(1.5)


def cell_shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    # shd phải đứng trước vAlign trong tcPr (theo schema OOXML)
    valign = tcpr.find(qn("w:vAlign"))
    if valign is not None:
        valign.addprevious(shd)
    else:
        tcpr.append(shd)


def row_props(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trpr.append(th)


# ------------------------------------------------------------------- các trang
def centered(doc, text, size, bold=False, before=0, after=0, italic=False):
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=before, after=after, spacing=1.0)
    add_rich(p, text, size=size, bold=bold, italic=italic)
    return p


def school_block(doc):
    centered(doc, "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH", 13, before=6)
    centered(doc, "TRƯỜNG ĐẠI HỌC KHOA HỌC TỰ NHIÊN", 14, bold=True)
    rule = doc.add_paragraph()
    para_fmt(rule, first=0, left=5.25, right=5.25, after=0, spacing=1.0)
    bottom_rule(rule)


def title_block(doc, before):
    for k, line in enumerate(nd.TEN_DE_TAI_DONG):
        centered(doc, line, 18, bold=True, before=before if k == 0 else 0, after=2)


def course_block(doc, before):
    centered(doc, "Tiểu luận Triết học", 14, bold=True, before=before)
    centered(doc, "Chương trình cao học và nghiên cứu sinh", 14, bold=True)
    centered(doc, "không chuyên ngành Triết học", 14, bold=True)


def cover_pages(doc):
    # Trang bìa
    school_block(doc)
    centered(doc, nd.HO_TEN, 16, bold=True, before=48)
    centered(doc, "MSHV: " + nd.MSHV, 14)
    title_block(doc, before=66)
    course_block(doc, before=24)
    centered(doc, "TP. HỒ CHÍ MINH – " + nd.NAM, 14, bold=True, before=210)
    # Trang áp bìa
    pb = doc.add_paragraph()
    para_fmt(pb, first=0, after=0, spacing=1.0)
    pb.add_run().add_break(WD_BREAK.PAGE)
    school_block(doc)
    title_block(doc, before=66)
    course_block(doc, before=24)
    for k, (txt, b) in enumerate(((nd.HO_TEN, True), ("MSHV: " + nd.MSHV, False),
                                  ("Ngành đang được đào tạo: " + nd.NGANH, False))):
        p = doc.add_paragraph()
        para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, left=7.0, before=60 if k == 0 else 0,
                 after=0, spacing=1.0)
        add_rich(p, txt, size=14, bold=b)
    centered(doc, "TP. HỒ CHÍ MINH – " + nd.NAM, 14, bold=True, before=180)


def toc_line(doc, text, page, level, bold=False, italic=False):
    indent = {1: 0.0, 2: 0.5, 3: 1.0}[level]
    hang = {1: 0.0, 2: 0.9, 3: 1.2}[level]
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, left=indent + hang, first=-hang, right=0.6,
             before=0, after=0, spacing=DAN_DONG)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(BE_RONG_CHU), WD_TAB_ALIGNMENT.RIGHT,
                                              WD_TAB_LEADER.DOTS)
    add_rich(p, text, size=CO_CHU, bold=bold, italic=italic)
    r = p.add_run("\t" + str(page))
    r.font.size = Pt(CO_CHU)
    r.font.bold = bold
    _set_fonts(r._r)
    return p


def to_roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def toc_entries():
    """Danh sách (cấp, nhãn trong mục lục, văn bản đề mục trên trang)."""
    out = []
    for block in nd.NOI_DUNG:
        if block[0] == "h1":
            out.append((1, block[2], block[1].replace("\n", " ")))
        elif block[0] in ("h2", "h3"):
            out.append((int(block[0][1]), block[1], block[1]))
    out.append((1, "DANH MỤC TÀI LIỆU THAM KHẢO", "DANH MỤC TÀI LIỆU THAM KHẢO"))
    return out


def front_matter(doc, pages):
    p = centered(doc, "MỤC LỤC", CO_CHUONG, bold=True, after=12)
    p.paragraph_format.line_spacing = DAN_DONG
    toc_line(doc, "DANH MỤC CHỮ VIẾT TẮT", pages.get("__cvt", "i"), 1, bold=True)
    toc_line(doc, "DANH MỤC BẢNG", pages.get("__bang", "i"), 1, bold=True)
    toc_line(doc, "DANH MỤC HÌNH", pages.get("__hinh", "i"), 1, bold=True)
    for k, (lvl, label, _) in enumerate(toc_entries()):
        toc_line(doc, label, pages.get(k, 0), lvl, bold=(lvl == 1), italic=(lvl == 3))

    # Danh mục chữ viết tắt (trang mới)
    p = centered(doc, "DANH MỤC CHỮ VIẾT TẮT", CO_CHUONG, bold=True, after=12)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.line_spacing = DAN_DONG
    for ab, full in nd.CHU_VIET_TAT:
        q = doc.add_paragraph()
        para_fmt(q, align=WD_ALIGN_PARAGRAPH.LEFT, left=2.7, first=-2.7, after=0, spacing=DAN_DONG)
        q.paragraph_format.tab_stops.add_tab_stop(Cm(2.7))
        add_rich(q, ab, size=CO_CHU, bold=True)
        add_rich(q, "\t" + full, size=CO_CHU)

    p = centered(doc, "DANH MỤC BẢNG", CO_CHUONG, bold=True, after=12)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.line_spacing = DAN_DONG
    for key in figure_table_order("table"):
        toc_line(doc, nd.BANG[key]["tieu_de"], pages.get(("bang", key), 0), 1)

    p = centered(doc, "DANH MỤC HÌNH", CO_CHUONG, bold=True, before=24, after=12)
    p.paragraph_format.line_spacing = DAN_DONG
    for key in figure_table_order("figure"):
        toc_line(doc, nd.HINH[key]["tieu_de"], pages.get(("hinh", key), 0), 1)


def figure_table_order(kind):
    return [b[1] for b in nd.NOI_DUNG if b[0] == kind]


def add_heading(doc, block, first_in_section):
    kind = block[0]
    if kind == "h1":
        p = doc.add_paragraph(style="Heading 1")
        lines = block[1].split("\n")
        for k, line in enumerate(lines):
            if k:
                p.add_run().add_break()
            add_rich(p, line)
        p.paragraph_format.page_break_before = not first_in_section
        return p
    p = doc.add_paragraph(style="Heading 2" if kind == "h2" else "Heading 3")
    add_rich(p, block[1])
    return p


def add_table(doc, key):
    spec = nd.BANG[key]
    cap = doc.add_paragraph()
    para_fmt(cap, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=6, after=6, keep_next=True)
    add_rich(cap, spec["tieu_de"], bold=True)
    widths = [Cm(w) for w in spec["rong"]]
    assert abs(sum(spec["rong"]) - BE_RONG_CHU) < 0.05, key
    t = doc.add_table(rows=len(spec["hang"]), cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for i, row in enumerate(spec["hang"]):
        row_props(t.rows[i], header=(i == 0))
        for j, txt in enumerate(row):
            c = t.cell(i, j)
            c.width = widths[j]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT,
                     first=0, before=2, after=2, spacing=1.0)
            lines = txt.split("\n")
            for k, line in enumerate(lines):
                if k:
                    p.add_run().add_break()
                add_rich(p, line, size=CO_CHU, bold=(i == 0 or (j == 0 and k == 0)))
            if i == 0:
                cell_shade(c, "D9D9D9")
    src = doc.add_paragraph()
    para_fmt(src, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, before=4, after=8)
    add_rich(src, spec["nguon"], italic=True)


def add_figure(doc, key):
    spec = nd.HINH[key]
    p = doc.add_paragraph()
    para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=6, after=4, spacing=1.0,
             keep_next=True)
    p.add_run().add_picture(os.path.join(HERE, spec["tep"]), width=Cm(spec["rong_cm"]))
    cap = doc.add_paragraph()
    para_fmt(cap, align=WD_ALIGN_PARAGRAPH.CENTER, first=0, before=2, after=8)
    add_rich(cap, spec["tieu_de"], bold=True)


def make_figures():
    """Vẽ Hình 4.1: các cấp độ tiếp cận bản chất trong đánh giá hoạt tính ức chế enzyme."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    plt.rcParams["font.family"] = "Liberation Serif"
    levels = [
        ("1. Tín hiệu sơ cấp", "Tín hiệu lặp lại, thuộc về chính hợp chất", "Nhiễu quang học; tạp chất"),
        ("2. Liều – đáp ứng", "Quy luật phụ thuộc nồng độ", "Đường cong dốc\nbất thường"),
        ("3. Tính đặc hiệu", "Tác dụng đặc hiệu với enzyme đích", "Kết tập keo; ion kim loại;\n"
                                                                  "phản ứng không đặc hiệu"),
        ("4. Cơ chế", "Kiểu ức chế, hằng số ức chế, liên kết trực tiếp", "Suy diễn cơ chế\n"
                                                                       "từ điểm số docking"),
        ("5. Quan hệ cấu trúc – hoạt tính", "Tính nhất quán trong dãy đồng loại", "SAR “phẳng”"),
        ("6. Tương ứng với đích sinh học", "Tác dụng trên đích ở người, tế bào, cơ thể",
         "Mô hình enzyme\nkhông tương đồng"),
    ]
    fig, ax = plt.subplots(figsize=(6.3, 6.0))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 106)
    ax.axis("off")
    h, gap, y0 = 12.2, 3.2, 6.0
    ax.text(37.5, 102.5, "Cấp độ tiếp cận bản chất", ha="center", va="center", fontsize=10.5,
            fontweight="bold")
    ax.text(84.0, 102.5, "Giả tượng cần loại trừ", ha="center", va="center", fontsize=10.5,
            fontweight="bold")
    for i, (name, q, fake) in enumerate(levels):
        y = y0 + i * (h + gap)
        shade = 0.95 - 0.07 * i
        ax.add_patch(FancyBboxPatch((10, y), 55, h, boxstyle="round,pad=0.3,rounding_size=1.2",
                                    fc=(shade, shade, shade), ec="black", lw=0.9))
        ax.text(37.5, y + h * 0.66, name, ha="center", va="center", fontsize=10, fontweight="bold")
        ax.text(37.5, y + h * 0.28, q, ha="center", va="center", fontsize=9)
        ax.add_patch(FancyBboxPatch((69, y), 30, h, boxstyle="round,pad=0.3,rounding_size=1.2",
                                    fc="white", ec="black", lw=0.8, ls="--"))
        ax.text(84, y + h / 2, fake, ha="center", va="center", fontsize=9, style="italic")
        ax.annotate("", xy=(69, y + h / 2), xytext=(65.6, y + h / 2),
                    arrowprops=dict(arrowstyle="-", lw=0.8, color="black", ls=":"))
        if i < len(levels) - 1:
            ax.annotate("", xy=(37.5, y + h + gap - 0.2), xytext=(37.5, y + h + 0.3),
                        arrowprops=dict(arrowstyle="-|>", lw=1.0, color="black"))
    top = y0 + len(levels) * (h + gap) - gap
    ax.annotate("", xy=(4, top), xytext=(4, y0), arrowprops=dict(arrowstyle="-|>", lw=1.6,
                                                                color="black"))
    ax.text(1.2, (top + y0) / 2, "Nhận thức đi sâu từ hiện tượng đến bản chất", rotation=90,
            ha="center", va="center", fontsize=9.5)
    ax.text(4, y0 - 3.2, "HIỆN TƯỢNG", ha="center", va="center", fontsize=8.5, fontweight="bold")
    ax.text(4, top + 3.0, "BẢN CHẤT", ha="center", va="center", fontsize=8.5, fontweight="bold")
    fig.savefig(os.path.join(HERE, nd.HINH["H4_1"]["tep"]), dpi=300, bbox_inches="tight")
    plt.close(fig)


def references(doc):
    p = doc.add_paragraph(style="Heading 1")
    add_rich(p, "DANH MỤC TÀI LIỆU THAM KHẢO")
    p.paragraph_format.page_break_before = True
    n = 0
    for label, items in (("Tài liệu tiếng Việt", nd.TAI_LIEU_VIET),
                         ("Tài liệu tiếng Anh", nd.TAI_LIEU_ANH)):
        q = doc.add_paragraph()
        para_fmt(q, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, before=6, after=3, keep_next=True)
        add_rich(q, label, bold=True)
        for _, ref in items:
            n += 1
            r = doc.add_paragraph()
            para_fmt(r, align=WD_ALIGN_PARAGRAPH.JUSTIFY, left=1.0, first=-1.0, after=4)
            r.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
            add_rich(r, "[%d]\t" % n + ref)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    style_font(normal, CO_CHU)
    pf = normal.paragraph_format
    pf.line_spacing = DAN_DONG
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.widow_control = True
    specs = {"Heading 1": (CO_CHUONG, True, False), "Heading 2": (CO_CHU, True, False),
             "Heading 3": (CO_CHU, True, True)}
    for name, (size, bold, italic) in specs.items():
        st = doc.styles[name]
        style_font(st, size, bold, italic)
        f = st.paragraph_format
        f.keep_with_next = True
        f.line_spacing = DAN_DONG
        f.first_line_indent = Cm(0)
        f.left_indent = Cm(0)
        if name == "Heading 1":
            f.alignment = WD_ALIGN_PARAGRAPH.CENTER
            f.space_before, f.space_after = Pt(0), Pt(12)
        else:
            f.alignment = WD_ALIGN_PARAGRAPH.LEFT
            f.space_before, f.space_after = Pt(6), Pt(3)


def build(pages, out_docx):
    doc = Document()
    setup_styles(doc)
    doc.core_properties.author = "Nguyễn Đoàn Diễm Ngọc"
    doc.core_properties.title = " ".join(nd.TEN_DE_TAI_DONG).capitalize()

    s1 = doc.sections[0]
    set_margins(s1)
    page_borders(s1)
    cover_pages(doc)

    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(s2)
    remove_page_borders(s2)
    page_numbering(s2, "lowerRoman", 1)
    header_page_field(s2)
    front_matter(doc, pages)

    s3 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(s3)
    remove_page_borders(s3)
    page_numbering(s3, "decimal", 1)
    header_page_field(s3)

    first = True
    for block in nd.NOI_DUNG:
        kind = block[0]
        if kind in ("h1", "h2", "h3"):
            add_heading(doc, block, first)
            first = False
        elif kind == "sub":
            p = doc.add_paragraph()
            para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, first=0, before=6, after=3, keep_next=True)
            add_rich(p, block[1], bold=True)
        elif kind == "p":
            p = doc.add_paragraph()
            para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.0)
            add_rich(p, block[1])
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "figure":
            add_figure(doc, block[1])
    thieu = [k for k in REF_KEYS if k not in CITED]
    if thieu:
        raise RuntimeError("Tài liệu tham khảo chưa được trích dẫn: " + ", ".join(thieu))
    references(doc)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.save(out_docx)


# ------------------------------------------------------------- xác định số trang
def to_pdf(docx_path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.splitext(docx_path)[0] + ".pdf"


def _norm(s):
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", s))


def page_texts(pdf_path):
    pdf = pdfium.PdfDocument(pdf_path)
    return [_norm(pdf[i].get_textpage().get_text_range()) for i in range(len(pdf))]


def locate(pdf_path):
    texts = page_texts(pdf_path)
    toc_start = next(i for i, t in enumerate(texts) if "MỤCLỤC" in t)
    body = next(i for i, t in enumerate(texts) if "1.Tínhcấpthiếtcủađềtài" in t)
    cvt = next(i for i in range(toc_start + 1, body) if "DANHMỤCCHỮVIẾTTẮT" in texts[i])
    pages = {"__cvt": to_roman(cvt - toc_start + 1)}
    pages["__bang"] = to_roman(next(i for i in range(cvt, body) if "DANHMỤCBẢNG" in texts[i])
                               - toc_start + 1)
    pages["__hinh"] = to_roman(next(i for i in range(cvt, body) if "DANHMỤCHÌNH" in texts[i])
                               - toc_start + 1)
    cur = body
    for k, (_, _, heading) in enumerate(toc_entries()):
        key = _norm(heading)[:45]
        while key not in texts[cur]:
            cur += 1
            if cur >= len(texts):
                raise RuntimeError("Không tìm thấy đề mục: " + heading)
        pages[k] = cur - body + 1
    for kind, store in (("table", "bang"), ("figure", "hinh")):
        for key in figure_table_order(kind):
            spec = nd.BANG[key] if kind == "table" else nd.HINH[key]
            cap = _norm(spec["tieu_de"])[:40]
            pages[(store, key)] = next(i for i in range(body, len(texts)) if cap in texts[i]) - body + 1
    return pages, len(texts), body


def main():
    out_docx = os.path.join(HERE, TEN_TEP + ".docx")
    make_figures()
    pages = {}
    for _ in range(3):
        build(pages, out_docx)
        pdf = to_pdf(out_docx)
        new_pages, n_pages, body = locate(pdf)
        if new_pages == pages:
            break
        pages = new_pages
    build(pages, out_docx)
    pdf = to_pdf(out_docx)
    final, n_pages, body = locate(pdf)
    assert final == pages, "Số trang trong mục lục chưa ổn định"
    ket_luan = max(v for k, v in pages.items() if isinstance(k, int)
                   and toc_entries()[k][1] == "KẾT LUẬN")
    print("Tổng số trang PDF:", n_pages, "| trang 1 (Mở đầu) ở trang PDF thứ", body + 1)
    print("Kết luận bắt đầu ở trang", ket_luan)
    for k, (lvl, label, _) in enumerate(toc_entries()):
        print("  " * (lvl - 1) + label[:70], "....", pages[k])


if __name__ == "__main__":
    main()
